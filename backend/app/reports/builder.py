"""Report generation (FR-REP-01..08): one content model rendered as PDF, Word and PowerPoint."""
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd

from .. import config
from ..defaults import PARAMS, diff_from_default
from ..models import Dataset, Insight, ParameterSet, Project, Run, User
from . import charts

SECTIONS = ["summary", "data", "method", "classification", "horizontal", "terminal", "cruise", "leveloffs",
            "insights", "annex"]
SECTION_TITLES = {"summary": "Executive summary", "data": "Data and coverage", "method": "Method and parameters",
                  "classification": "Route classification", "horizontal": "Horizontal efficiency",
                  "terminal": "Terminal arrival", "cruise": "Cruise vertical efficiency",
                  "leveloffs": "Climb and descent level-offs", "insights": "Insights and recommendations",
                  "annex": "Annex"}


def fmt(v, d=1, unit=""):
    if v is None:
        return "–"
    try:
        return f"{v:,.{d}f}{unit}"
    except (TypeError, ValueError):
        return str(v)


def pct(v, d=0):
    return "–" if v is None else f"{100 * v:.{d}f}%"


def build_content(db, run: Run, opts: dict, out_dir: Path) -> dict:
    project: Project = db.get(Project, run.project_id)
    ds: Dataset = db.get(Dataset, run.dataset_id)
    ps: ParameterSet = db.get(ParameterSet, run.paramset_id)
    params = ps.values
    rd = config.run_dir(run.id)
    from ..engine import aggregate
    S = aggregate.compute(rd, params, project, opts.get("filters") or {}, ds.quality)
    meta = json.loads((rd / "meta.json").read_text())
    res = pd.read_parquet(rd / "flights.parquet")
    res = aggregate.apply_filters(res, opts.get("filters") or {})
    sections = [s for s in SECTIONS if s in (opts.get("sections") or SECTIONS)]
    ins = db.query(Insight).filter_by(run_id=run.id).order_by(Insight.effect.desc()).all()
    keep = [i for i in ins if i.status in ("accepted", "edited") or (opts.get("include_drafts") and i.status == "drafted")]
    insights = [i for i in keep if i.kind == "insight"]
    recs = [i for i in keep if i.kind == "recommendation"]
    B = []  # blocks
    H = lambda t, l=1: B.append({"t": "h", "level": l, "text": t})  # noqa: E731
    Pp = lambda t: B.append({"t": "p", "text": t})  # noqa: E731
    L = lambda items: B.append({"t": "ul", "items": items})  # noqa: E731
    T = lambda head, rows, widths=None: B.append({"t": "table", "head": head, "rows": rows, "widths": widths})  # noqa: E731
    I = lambda p, cap: p and B.append({"t": "img", "path": str(p), "caption": cap})  # noqa: E731
    PB = lambda: B.append({"t": "pb"})  # noqa: E731
    figs = {}
    main = S["main_ring"]
    # charts
    from ..engine.pipeline import load_navdb
    fx, aw = load_navdb(db)
    import numpy as np
    import pyarrow.dataset as pads
    fids = res.fid.tolist()
    pts = pads.dataset(config.dataset_dir(run.dataset_id) / "points.parquet").to_table(
        columns=["fid", "t", "lat", "lon", "alt"], filter=pads.field("fid").isin(fids)).to_pandas()
    tracks = []
    meta_r = res.set_index("fid")
    for fid, g in pts.groupby("fid", sort=False):
        g = g.sort_values("t")
        g = g[g.alt > 0] if (g.alt > 0).sum() > 5 else g
        idx = np.linspace(0, len(g) - 1, min(len(g), 50)).astype(int)
        tracks.append({"dirn": meta_r.loc[fid].dirn, "outlier": bool(meta_r.loc[fid].outlier),
                       "coords": np.c_[g.lat.values[idx], g.lon.values[idx]].tolist()})
    routes = {}
    from ..engine.fpl import expand_route
    for d, g in res.groupby("dirn"):
        rk = g.route_key.dropna()
        if len(rk):
            top = Counter(rk).most_common(1)[0][0]
            names, _, _ = expand_route(top, dict(fx), aw)
            a, b = d.split("-")
            ap = {project.adep: (project.adep_lat, project.adep_lon), project.ades: (project.ades_lat, project.ades_lon)}
            coords = [ap[a]] + [fx[n] for n in names if n in fx] + [ap[b]]
            routes[d] = [(top, coords)]
    ap = {project.adep: (project.adep_lat, project.adep_lon), project.ades: (project.ades_lat, project.ades_lon)}
    figs["map"] = charts.track_map(tracks, routes, ap, out_dir / "f_map.png") if tracks else None
    figs["ext"] = charts.extension_rings(S, out_dir / "f_ext.png") if S["extension"] else None
    figs["enr"] = charts.enroute_monthly(S, out_dir / "f_enr.png") if S["enroute_monthly"] else None
    figs["lat"] = charts.lateral(S, out_dir / "f_lat.png") if S["lateral"] else None
    figs["ter"] = charts.terminal_monthly(S, out_dir / "f_ter.png") if S["terminal"]["monthly"] else None
    figs["cru"] = charts.cruise_classes(S, out_dir / "f_cru.png")
    figs["los"] = charts.leveloff_sensitivity(S, out_dir / "f_los.png") if S["leveloffs"]["sensitivity"] else None
    figs["loh"] = charts.leveloff_hist(S, out_dir / "f_loh.png", "descent")
    figs["loc"] = charts.leveloff_hist(S, out_dir / "f_loc.png", "climb")
    figs["clu"] = charts.cluster_monthly(S, out_dir / "f_clu.png")

    title = opts.get("title") or f"{project.adep}–{project.ades} trajectory efficiency analysis"
    authors = opts.get("authors") or ""
    kpis = S["kpis"]

    if "summary" in sections:
        H(SECTION_TITLES["summary"])
        top = (insights or [])[:5]
        if top:
            Pp(top[0].edited or top[0].drafted)
            L([(i.title + ": " if i.title else "") + (i.edited or i.drafted) for i in top[1:]])
        else:
            Pp("No insights have been accepted yet for this run.")
        T(["Indicator", "Value"], [[k["label"], f"{k['value']} {k['unit']}".strip()] for k in kpis], [11, 5])
    if "data" in sections:
        H(SECTION_TITLES["data"])
        q = ds.quality or {}
        Pp(f"The analysis covers {S['n']['flights']} flights of {project.adep}–{project.ades} "
           f"({', '.join(f'{k}: {v}' for k, v in S['n']['by_dirn'].items())}), from dataset “{ds.name}” "
           f"(version {ds.version}, {ds.n_flights} flights loaded). {S['n']['with_plan']} flights are matched to a flight plan "
           f"and {S['n']['with_ref']} have a fully expanded reference route.")
        cov = q.get("coverage") or {}
        if cov.get("months"):
            months = cov["months"]
            T(["Callsign"] + [m[2:] for m in months], [[cs] + [str(x) for x in row] for cs, row in zip(cov["callsigns"], cov["counts"])])
            B[-1]["caption"] = "Flights per callsign per month"
        g = q.get("gaps") or {}
        L([f"Gate-to-gate tracks: {pct(q.get('gate_to_gate_share'), 1)}",
           "Sampling interval (median): " + "; ".join(f"{b['band']} {fmt(b['median_s'], 0)} s" for b in q.get("sampling", [])),
           f"Airborne gaps longer than {g.get('threshold_s')} s: {g.get('n')} in {g.get('flights')} flights; longer than "
           f"{g.get('major_threshold_s')} s: {g.get('n_major')} in {g.get('flights_major')} flights",
           "Aircraft types (Item 9): " + ", ".join(f"{t['callsign']} {t['type']} ({t['n']})" for t in S["types"][:12])])
        I(figs["map"], "Figure: tracks and the most frequent filed route per direction")
    if "method" in sections:
        H(SECTION_TITLES["method"])
        Pp("Method: ENRI End-to-End Analysis Methodology v0.3 with CAAS comments, and ENRI's leg-based cruise method "
           "(13 October 2026). Every metric is computed from the run's stored dataset and parameter set.")
        d = diff_from_default(params)
        L([f"Parameter set: {ps.name}, version {ps.version}",
           f"Earth model: {params['GEN-01']} (R = 3,440.065 NM)" if params["GEN-01"] == "spherical" else "Earth model: WGS84 ellipsoid",
           f"En-route boundary: {main} NM rings; sensitivity {', '.join(str(int(x)) for x in params['HOR-01b'])} NM",
           f"Level-off: |VR| ≤ {params['VER-02']} ft/s for ≥ {params['VER-03']} s; climb bounds {params['VER-07']}, descent bounds {params['VER-08']}",
           f"Values differing from the VTBS–WADD default: {', '.join(f'{k} = {v}' for k, v in d.items()) or 'none'}"])
    if "classification" in sections and S.get("clusters"):
        H(SECTION_TITLES["classification"])
        T(["Direction", "Cluster", "Flights", "Share", "Outliers", "Contingency"],
          [[c["dirn"], str(c["cluster"]), str(c["n"]), pct(c["share"]), str(c["outliers"]), "yes" if c["contingency"] else "no"]
           for c in S["clusters"]["summary"]])
        I(figs["clu"], "Figure: flights per route cluster by month (metric 7)")
    if "horizontal" in sections:
        H(SECTION_TITLES["horizontal"])
        T(["Ring", "Direction", "n", "Reference (NM)", "Flown (NM)", "Extension median (NM)", "P15 / P85", "Shorter than filed",
           "En-route time (min)", "Time variance (min)"],
          [[str(e["ring"]), e["dirn"], str(e["n"]), fmt(e["ref_median"]), fmt(e["flown_median"]), fmt(e["median"]),
            f"{fmt(e['p15'])} / {fmt(e['p85'])}", pct(e["shorter_share"]), fmt(e["time_median"]), fmt(e["time_var"])]
           for e in S["extension"]])
        I(figs["ext"], "Figure: path extension against the filed route by boundary radius (metrics 1–3)")
        I(figs["enr"], "Figure: en-route time by month (metrics 4–5)")
        I(figs["lat"], "Figure: lateral deviation from the filed route")
        if S["fir"]:
            T(["Direction", "FIR", "ANSP", "Flights", "Flown (NM)", "Reference (NM)", "Extension (NM)", "Time (min)"],
              [[f["dirn"], f["fir"], f["ansp"], str(f["n"]), fmt(f["flown_median"]), fmt(f["ref_median"]), fmt(f["ext_median"]),
                fmt(f["time_median"])] for f in S["fir"]])
        else:
            Pp("Airspace (FIR) efficiency: no FIR boundaries are loaded, so metric 6 is not reported.")
    if "terminal" in sections:
        H(SECTION_TITLES["terminal"])
        T(["Destination", "Ring (NM)", "n", "Median (min)", "P85 (min)", "Variance (min)", "Unimpeded (min)", "Extra (min)",
           "Holding-type"],
          [[t["ades"], str(t["ring"]), str(t["n"]), fmt(t["median"]), fmt(t["p85"]), fmt(t["var"]), fmt(t["unimpeded"]),
            fmt(t["add_mean"]), pct(t["hold_share"])] for t in S["terminal"]["by_ades"]])
        corr = S["terminal"]["corr"]
        if corr:
            Pp("Correlation of en-route time with arrival-ring time: " +
               ", ".join(f"{k} r = {fmt(v, 2)}" for k, v in corr.items()) + ".")
        I(figs["ter"], "Figure: monthly extra arrival time and holding share")
    if "cruise" in sections:
        H(SECTION_TITLES["cruise"])
        cr = S["cruise"]
        if cr["score"]:
            T(["Direction", "Type", "Flights", "Score median", "Score P85", "Time above (min)", "Time below (min)", "Time at (min)"],
              [[c["dirn"], c["type"], str(c["n"]), fmt(c["median"]), fmt(c["p85"]), fmt(c["time_above_mean"]),
                fmt(c["time_below_mean"]), fmt(c["time_at_mean"])] for c in cr["score"]])
            B[-1]["caption"] = f"Metric 8, leg-based score (penalty: {params['CRU-07']}; unit = 1,000 ft × min)"
        I(figs["cru"], "Figure: flown cruise level against filed level at flight-plan waypoints")
        ch = cr.get("changes") or {}
        if ch:
            L([f"Flown cruise level changes: {ch['n']} ({pct(ch['climb_share'])} climbs), median transition {fmt(ch['transition_median_min'])} min",
               f"Changes reaching a filed level within {params['CRU-08']:.0f} NM of its filed waypoint: {pct(ch.get('matched_share'))} "
               f"(median distance {fmt(ch.get('median_dist_planned'), 0)} NM)",
               f"Changes on flights that filed no level change: {pct(ch['on_unplanned_share'])}"])
        if cr["capped"]:
            T(["Direction", "Type", "Flights", "Capped climb", "Median below filed (ft)", "Median distance (NM)", "Median duration (min)"],
              [[c["dirn"], c["type"], str(c["n"]), pct(c["share"]), fmt(c["median_below_ft"], 0), fmt(c["median_dist"], 0),
                fmt(c["median_min"])] for c in cr["capped"]])
    if "leveloffs" in sections:
        H(SECTION_TITLES["leveloffs"])
        T(["Phase", "Airport", "Flights", "With level-off", "Per flight", "Median duration (s)", "Time level per flight (min)"],
          [[p["phase"], p["airport"], str(p["flights"]), pct(p["share"]), fmt(p["mean_per_flight"], 2), fmt(p["median_dur"], 0),
            fmt(p["time_per_flight_min"], 2)] for p in S["leveloffs"]["primary"]])
        B[-1]["caption"] = (f"Level-off rule: |VR| ≤ {params['VER-02']} ft/s for ≥ {params['VER-03']} s; "
                            f"final-approach intercept below {params['VER-10'][0]:.0f} ft within {params['VER-10'][1]:.0f} NM excluded")
        I(figs["los"], "Figure: share of flights with a level-off by minimum duration (metrics 9–10)")
        I(figs["loh"], "Figure: altitude of descent level segments ≥ 60 s")
        I(figs["loc"], "Figure: altitude of climb level segments ≥ 60 s")
        T(["Phase", "Airport", "Band", "Flights", "Median duration (s)", "Median distance (NM)", "Min per flight"],
          [[b["phase"], b["airport"], b["band"], str(b["flights"]), fmt(b["median_dur"], 0), fmt(b["median_dist"], 0),
            fmt(b["min_per_flight"], 2)] for b in S["leveloffs"]["bands"] if b["flights"]])
    if "insights" in sections:
        H(SECTION_TITLES["insights"])
        if insights:
            for i in insights:
                B.append({"t": "insight", "title": i.title, "text": i.edited or i.drafted, "rule": i.rule})
        else:
            Pp("No accepted insights.")
        if recs:
            H("Recommendations", 2)
            L([(r.title + ": " if r.title else "") + (r.edited or r.drafted) for r in recs])
    if "annex" in sections:
        H(SECTION_TITLES["annex"])
        H("Parameter set", 2)
        T(["ID", "Parameter", "Value", "Unit", "Source"],
          [[k, PARAMS[k]["label"], json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else str(v),
            PARAMS[k]["unit"], PARAMS[k]["source"]] for k, v in params.items() if k in PARAMS and k != "REP-02"])
        ex = meta.get("excluded") or []
        H("Excluded flights", 2)
        if ex:
            T(["Flight", "Reason"], [[e["fid"], e["exclusion"]] for e in ex[:200]])
        else:
            Pp("None.")
        uf = meta.get("unresolved_fixes") or {}
        if uf:
            H("Unresolved flight plan fixes", 2)
            T(["Fix / route element", "Flights"], [[k, str(v)] for k, v in sorted(uf.items(), key=lambda x: -x[1])])
        H("Glossary", 2)
        T(["Term", "Meaning"], [["ADEP / ADES", "Departure / destination aerodrome"],
                                ["Reference route", "The filed route expanded into waypoint coordinates"],
                                ["En-route portion", "The part of a flight between the boundary rings"],
                                ["Level segment", "A period with |vertical rate| at or below the level threshold"],
                                ["TOC / TOD", "Top of climb / top of descent"],
                                ["RFL", "Requested (filed) flight level"],
                                ["CPA", "Closest point of approach to a waypoint"]])
    return {"title": title, "authors": authors, "blocks": B, "project": project, "run": run, "ds": ds, "ps": ps,
            "kpis": kpis, "insights": insights, "recs": recs, "figs": figs, "S": S,
            "stamp": f"Dataset {ds.name} v{ds.version} ({ds.content_hash[:10]}) · Parameters {ps.name} v{ps.version} · "
                     f"Earth model {params['GEN-01']} · CPTEA {run.software_version or config.VERSION} · Run {run.id} "
                     f"({run.finished_at:%Y-%m-%d %H:%M} UTC)" if run.finished_at else ""}


# ------------------------------------------------------------------ PDF
def render_pdf(c: dict, path: Path, logo: Path | None = None):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import matplotlib
    fdir = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
    try:
        pdfmetrics.registerFont(TTFont("DV", str(fdir / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DVB", str(fdir / "DejaVuSans-Bold.ttf")))
        pdfmetrics.registerFont(TTFont("DVI", str(fdir / "DejaVuSans-Oblique.ttf")))
        from reportlab.pdfbase.pdfmetrics import registerFontFamily
        registerFontFamily("DV", normal="DV", bold="DVB", italic="DVI", boldItalic="DVB")
        F, FB, FI = "DV", "DVB", "DVI"
    except Exception:  # noqa: BLE001
        F, FB, FI = "Helvetica", "Helvetica-Bold", "Helvetica-Oblique"
    ink, muted, acc, rule, tint = colors.HexColor("#1B1F24"), colors.HexColor("#6B7A82"), colors.HexColor("#0F7482"), \
        colors.HexColor("#D9DEE0"), colors.HexColor("#EEF3F4")
    st = {"p": ParagraphStyle("p", fontName=F, fontSize=9.2, leading=13, textColor=ink, spaceAfter=5),
          "h1": ParagraphStyle("h1", fontName=FB, fontSize=14, leading=18, textColor=ink, spaceBefore=10, spaceAfter=6),
          "h2": ParagraphStyle("h2", fontName=FB, fontSize=11, leading=14, textColor=acc, spaceBefore=8, spaceAfter=4),
          "cap": ParagraphStyle("cap", fontName=FI, fontSize=8, leading=10.5, textColor=muted, spaceAfter=9),
          "tc": ParagraphStyle("tc", fontName=F, fontSize=7.6, leading=9.6, textColor=ink),
          "th": ParagraphStyle("th", fontName=FB, fontSize=7.6, leading=9.6, textColor=ink),
          "title": ParagraphStyle("t", fontName=FB, fontSize=20, leading=25, textColor=ink, spaceAfter=8),
          "sub": ParagraphStyle("s", fontName=F, fontSize=10.5, leading=14, textColor=muted, spaceAfter=3),
          "ul": ParagraphStyle("ul", fontName=F, fontSize=9.2, leading=13, textColor=ink, leftIndent=12, bulletIndent=2, spaceAfter=2)}
    W = A4[0] - 36 * mm

    def esc(s):
        return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    story = []
    if logo and Path(logo).exists():
        story.append(Image(str(logo), width=30 * mm, height=15 * mm, kind="proportional"))
    story += [Paragraph(esc(c["title"]), st["title"]),
              Paragraph(esc(f"{c['project'].name} · {datetime.now(timezone.utc):%d %B %Y}"), st["sub"])]
    if c["authors"]:
        story.append(Paragraph(esc(c["authors"]), st["sub"]))
    story.append(Paragraph(esc(c["stamp"]), st["cap"]))
    story.append(Spacer(1, 6))
    for b in c["blocks"]:
        t = b["t"]
        if t == "h":
            story.append(Paragraph(esc(b["text"]), st["h1" if b["level"] == 1 else "h2"]))
        elif t == "p":
            story.append(Paragraph(esc(b["text"]), st["p"]))
        elif t == "ul":
            for it in b["items"]:
                story.append(Paragraph(esc(it), st["ul"], bulletText="•"))
            story.append(Spacer(1, 4))
        elif t == "insight":
            inner = [Paragraph(f"<b>{esc(b['title'])}</b>" if b["title"] else "", st["p"]), Paragraph(esc(b["text"]), st["p"])]
            tb = Table([[inner]], colWidths=[W])
            tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), tint), ("LEFTPADDING", (0, 0), (-1, -1), 8),
                                    ("RIGHTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 5),
                                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
            story += [tb, Spacer(1, 5)]
        elif t == "table":
            n = len(b["head"])
            ws = b.get("widths")
            cw = [W * w / sum(ws) for w in ws] if ws else [W / n] * n
            data = [[Paragraph(esc(h), st["th"]) for h in b["head"]]] + [[Paragraph(esc(x), st["tc"]) for x in r] for r in b["rows"]]
            tb = Table(data, colWidths=cw, repeatRows=1)
            tb.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, -1), .4, rule),
                                    ("BACKGROUND", (0, 0), (-1, 0), tint), ("LINEBELOW", (0, 0), (-1, 0), .8, acc),
                                    ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
                                    ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3)]))
            story.append(tb)
            story.append(Paragraph(esc(b["caption"]), st["cap"]) if b.get("caption") else Spacer(1, 8))
        elif t == "img":
            from PIL import Image as PI
            im = PI.open(b["path"])
            w = W * 0.95
            h = w * im.size[1] / im.size[0]
            if h > 150 * mm:
                h = 150 * mm
                w = h * im.size[0] / im.size[1]
            story.append(KeepTogether([Image(b["path"], width=w, height=h), Paragraph(esc(b["caption"]), st["cap"])]))
        elif t == "pb":
            story.append(PageBreak())

    def foot(cv, doc):
        cv.saveState()
        cv.setFont(F, 7.5)
        cv.setFillColor(muted)
        cv.drawString(18 * mm, 10 * mm, f"CPTEA · {c['project'].name} · Run {c['run'].id}")
        cv.drawRightString(A4[0] - 18 * mm, 10 * mm, str(doc.page))
        cv.restoreState()

    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm,
                            bottomMargin=16 * mm, title=c["title"], author=c["authors"] or "CPTEA")
    doc.build(story, onFirstPage=foot, onLaterPages=foot)
    return path


# ------------------------------------------------------------------ Word
def render_docx(c: dict, path: Path, logo: Path | None = None):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Cm, Pt, RGBColor
    d = Document()
    stl = d.styles["Normal"]
    stl.font.name = "Calibri"
    stl.font.size = Pt(10)
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(1.8)
    if logo and Path(logo).exists():
        d.add_picture(str(logo), width=Cm(3))
    d.add_heading(c["title"], 0)
    p = d.add_paragraph(f"{c['project'].name} · {datetime.now(timezone.utc):%d %B %Y}")
    if c["authors"]:
        d.add_paragraph(c["authors"])
    sp = d.add_paragraph(c["stamp"])
    sp.runs[0].font.size = Pt(8)
    sp.runs[0].font.color.rgb = RGBColor(0x6B, 0x7A, 0x82)
    for b in c["blocks"]:
        t = b["t"]
        if t == "h":
            d.add_heading(b["text"], b["level"])
        elif t == "p":
            d.add_paragraph(b["text"])
        elif t == "ul":
            for it in b["items"]:
                d.add_paragraph(it, style="List Bullet")
        elif t == "insight":
            p = d.add_paragraph()
            if b["title"]:
                p.add_run(b["title"] + ". ").bold = True
            p.add_run(b["text"])
        elif t == "table":
            tb = d.add_table(rows=1, cols=len(b["head"]))
            tb.style = "Light Grid Accent 1"
            tb.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, h in enumerate(b["head"]):
                tb.rows[0].cells[i].text = h
            for r in b["rows"]:
                cells = tb.add_row().cells
                for i, x in enumerate(r):
                    cells[i].text = str(x)
            for row in tb.rows:
                for cell in row.cells:
                    for pp in cell.paragraphs:
                        for run in pp.runs:
                            run.font.size = Pt(8)
            if b.get("caption"):
                cp = d.add_paragraph(b["caption"])
                cp.runs[0].italic = True
                cp.runs[0].font.size = Pt(8)
        elif t == "img":
            d.add_picture(b["path"], width=Cm(16))
            cp = d.add_paragraph(b["caption"])
            cp.runs[0].italic = True
            cp.runs[0].font.size = Pt(8)
        elif t == "pb":
            d.add_page_break()
    d.save(str(path))
    return path


# ------------------------------------------------------------------ PowerPoint summary deck
def render_pptx(c: dict, path: Path, logo: Path | None = None):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    from PIL import Image as PI
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    DARK, ACC, INK, MUT, LIGHT = RGBColor(0x0E, 0x3A, 0x47), RGBColor(0xE3, 0x9B, 0x2D), RGBColor(0x1B, 0x1F, 0x24), \
        RGBColor(0x6B, 0x7A, 0x82), RGBColor(0xEE, 0xF3, 0xF4)

    def text(s, x, y, w, h, t, size=16, bold=False, color=INK, font="Calibri"):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        lines = t if isinstance(t, list) else [t]
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = ln
            p.font.size = Pt(size)
            p.font.bold = bold
            p.font.color.rgb = color
            p.font.name = font
            p.space_after = Pt(6)
        return tb

    def bg(s, color):
        f = s.background.fill
        f.solid()
        f.fore_color.rgb = color

    s = prs.slides.add_slide(blank)
    bg(s, DARK)
    text(s, 0.6, 2.2, 12, 1.2, c["title"], 36, True, RGBColor(0xFF, 0xFF, 0xFF), "Cambria")
    text(s, 0.6, 3.5, 12, 0.6, f"{c['project'].name} · {datetime.now(timezone.utc):%d %B %Y}", 18, False, RGBColor(0x5F, 0xB3, 0xBE))
    text(s, 0.6, 6.6, 12, 0.5, c["stamp"], 10, False, RGBColor(0x5F, 0xB3, 0xBE))
    if logo and Path(logo).exists():
        s.shapes.add_picture(str(logo), Inches(11.3), Inches(0.4), height=Inches(0.9))
    # KPIs
    s = prs.slides.add_slide(blank)
    text(s, 0.6, 0.35, 12, 0.8, "Key indicators", 30, True, DARK, "Cambria")
    for i, k in enumerate(c["kpis"][:8]):
        col, row = i % 4, i // 4
        x, y = 0.6 + col * 3.1, 1.5 + row * 2.6
        sh = s.shapes.add_shape(5, Inches(x), Inches(y), Inches(2.9), Inches(2.3))
        sh.fill.solid()
        sh.fill.fore_color.rgb = LIGHT
        sh.line.fill.background()
        text(s, x + 0.2, y + 0.2, 2.5, 0.9, f"{k['value']}", 30, True, RGBColor(0x0F, 0x74, 0x82), "Cambria")
        text(s, x + 0.2, y + 1.1, 2.5, 1.1, [k["unit"], k["label"]], 12, False, INK)
    # chart slides
    order = [("map", "Tracks and filed routes"), ("ext", "Path extension against the filed route"),
             ("ter", "Terminal arrival"), ("cru", "Cruise level against filed"), ("los", "Level-off threshold sensitivity"),
             ("loh", "Where descent level-offs happen")]
    ins_by_chart = {}
    for i in c["insights"]:
        ins_by_chart.setdefault(i.chart, []).append(i)
    chart_map = {"ext": ["extension", "extension_rings"], "ter": ["terminal", "terminal_scatter", "terminal_monthly"],
                 "cru": ["cruise", "cruise_changes"], "los": ["leveloff_sensitivity"], "loh": ["leveloff_hist"], "map": ["clusters", "coverage"]}
    for key, ttl in order:
        fp = c["figs"].get(key)
        if not fp:
            continue
        s = prs.slides.add_slide(blank)
        text(s, 0.6, 0.35, 12, 0.8, ttl, 28, True, DARK, "Cambria")
        im = PI.open(fp)
        maxw, maxh = 8.2, 5.6
        w = maxw
        h = w * im.size[1] / im.size[0]
        if h > maxh:
            h = maxh
            w = h * im.size[0] / im.size[1]
        s.shapes.add_picture(str(fp), Inches(0.5), Inches(1.4), Inches(w), Inches(h))
        notes = [x for ch in chart_map.get(key, []) for x in ins_by_chart.get(ch, [])][:3]
        if notes:
            text(s, 9.0, 1.5, 3.9, 5.4, [(n.edited or n.drafted) for n in notes], 13, False, INK)
    # insights
    if c["insights"]:
        s = prs.slides.add_slide(blank)
        text(s, 0.6, 0.35, 12, 0.8, "Insights", 30, True, DARK, "Cambria")
        text(s, 0.6, 1.3, 12.1, 5.8, ["• " + (i.title or (i.edited or i.drafted)[:90]) for i in c["insights"][:10]], 15)
    if c["recs"]:
        s = prs.slides.add_slide(blank)
        bg(s, DARK)
        text(s, 0.6, 0.35, 12, 0.8, "Recommendations", 30, True, RGBColor(0xFF, 0xFF, 0xFF), "Cambria")
        text(s, 0.6, 1.4, 12.1, 5.6, [f"{k + 1}. " + (r.title + ": " if r.title else "") + (r.edited or r.drafted)
                                      for k, r in enumerate(c["recs"][:7])], 16, False, RGBColor(0xFF, 0xFF, 0xFF))
    prs.save(str(path))
    return path


def generate(db, report, run: Run) -> dict:
    out_dir = config.run_dir(run.id) / f"report_{report.id}"
    out_dir.mkdir(parents=True, exist_ok=True)
    opts = report.options or {}
    c = build_content(db, run, opts, out_dir)
    logo = Path(opts["logo_path"]) if opts.get("logo_path") else None
    base = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in c["title"])[:60] or "report"
    files = {}
    fmts = opts.get("formats") or ["pdf", "docx", "pptx"]
    if "pdf" in fmts:
        files["pdf"] = str(render_pdf(c, out_dir / f"{base}.pdf", logo))
    if "docx" in fmts:
        files["docx"] = str(render_docx(c, out_dir / f"{base}.docx", logo))
    if "pptx" in fmts:
        files["pptx"] = str(render_pptx(c, out_dir / f"{base}.pptx", logo))
    return files
